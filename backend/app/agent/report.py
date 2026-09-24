import json

def markdown(job):
    e=job['event'];c=job.get('correlation',{});lines=[f"# 调查报告：{e['name']}",'',f"调查 ID：{job['id']}",f"状态：{job['status']}",f"数据来源：{e['source']}（demo 为合成演示数据）",f"代码快照 SHA-256：{job.get('snapshot','未生成')}",'','## 结论','',c.get('summary','调查尚未完成'),'','所有关联均为待验证推断；不证明此次利用成功。','']
    for m in c.get('matches',[]):
        f=m['finding'];lines += [f"## {f['type']} · {f['cwe']}",'',f"入口：{m['entry']['method']} {m['entry']['path']}",f"漏洞位置：{f['file']}:{f['line']} / {f['function']}",f"关联置信度：{m['confidence']}（启发式等级，非统计概率）",'','调用链：'+' → '.join(x['name'] for x in m['chain']),'','### 依据','']+['- '+x for x in m['reasons']]+['','证据：'+', '.join(m['evidence_ids']),'','```python',f['code'],'```','','### 修复建议','',f['fix'],'','```python',f['patch'],'```','']
    lines += ['## 辅助研判（AI Inference）','',job.get('analysis',{}).get('text','未生成'),'','## 限制','']+['- '+x for x in job.get('review',{}).get('limitations',[])]+['- '+x for x in job.get('warnings',[])]+['','## 证据索引','']
    for v in job.get('evidence',[]): lines += [f"### {v['id']} · {v['kind']} · {v['title']}",f"来源：{v['source']} / {v['provenance']}",f"SHA-256：{v['sha256']}",'','```json',json.dumps(v['content'],ensure_ascii=False,indent=2),'```','']
    return '\n'.join(lines)

def word(job):
    """Actual OOXML document, with observed execution trace rather than hidden reasoning."""
    from io import BytesIO
    from docx import Document
    from docx.shared import Pt
    from docx.oxml.ns import qn
    doc=Document();style=doc.styles['Normal'];style.font.name='Microsoft YaHei';style.font.size=Pt(10)
    style.element.rPr.rFonts.set(qn('w:eastAsia'),'Microsoft YaHei')
    doc.add_heading('Double Pupil · 安全调查报告',0)
    doc.add_heading(job['event']['name'],1)
    for key,value in [('调查 ID',job['id']),('事件来源',job['event']['source']),('资产',job['event'].get('asset_ip','')),('代码项目',job['project']['name']),('代码快照 SHA-256',job.get('snapshot',''))]:doc.add_paragraph(f'{key}：{value}')
    if job.get('target_ips'):doc.add_paragraph('告警目的地址：'+', '.join(job['target_ips']))
    if job.get('target_assets'):doc.add_paragraph('目的资产：'+', '.join(a['name'] for a in job['target_assets']))
    doc.add_heading('调查结论',1);doc.add_paragraph(job.get('correlation',{}).get('summary','尚无结论'))
    doc.add_paragraph('静态关联是漏洞候选定位，不等于此次利用成功；生产部署版本仍需核验。')
    doc.add_heading('调查思维链 · 可审计执行过程',1)
    table=doc.add_table(rows=1,cols=4);table.style='Light Shading Accent 1'
    for cell,title in zip(table.rows[0].cells,['节点','工具','状态','动作摘要']):cell.text=title
    for step in job.get('timeline',[]):
        for cell,value in zip(table.add_row().cells,[step.get('name',''),step.get('tool',''),step.get('status',''),step.get('note','')]):cell.text=str(value)
    for match in job.get('correlation',{}).get('matches',[]):
        f=match['finding'];doc.add_heading(f"{f['cwe']} · {f['file']}:{f['line']}",1)
        doc.add_paragraph('攻击入口：'+str(match['entry'].get('method') or 'HTTP')+' '+match['entry']['path'])
        doc.add_paragraph('关联等级：'+match['confidence'])
        doc.add_paragraph(' → '.join(n['name'] for n in match['chain']))
        for reason in match['reasons']:doc.add_paragraph(reason,style='List Bullet')
        doc.add_heading('代码快照',2);doc.add_paragraph(f['code'])
        doc.add_heading('修复建议',2);doc.add_paragraph(f['fix'])
        if f.get('patch'):doc.add_paragraph(f['patch'])
        doc.add_paragraph('证据 ID：'+', '.join(match['evidence_ids']))
    doc.add_heading('模型辅助研判（未验证推断）',1);doc.add_paragraph(job.get('analysis',{}).get('text','未生成'))
    doc.add_heading('限制与告警',1)
    for limit in job.get('review',{}).get('limitations',[])+job.get('warnings',[]):doc.add_paragraph(limit,style='List Bullet')
    doc.add_heading('证据索引与来源',1)
    for evidence in job.get('evidence',[]):
        doc.add_heading(evidence['id']+' · '+evidence['title'],2)
        doc.add_paragraph(f"{evidence['kind']} / {evidence['source']} / {evidence['provenance']}")
        doc.add_paragraph('SHA-256：'+evidence['sha256'])
        raw=json.dumps(evidence['content'],ensure_ascii=False,indent=2)
        doc.add_paragraph(raw[:12000]+('\n[内容过长，完整内容见 JSON 报告]' if len(raw)>12000 else ''))
    output=BytesIO();doc.save(output);return output.getvalue()
